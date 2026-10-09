import asyncio
import os
import pandas as pd
import base64
import docx
import json
from pathlib import Path
import logging
from langchain_community.document_loaders import PyPDFLoader
import operator
from pymongo import AsyncMongoClient
import rapidfuzz

from get_llm_response import get_llm, get_answer, get_answer_with_schema
from exception_type import AIERROR

from ai_chatbot.src.configs.app import settings
from ai_chatbot.src.schemas.writer import SummaryStatus, SummaryAction, IntentStatus, Intent
from ai_chatbot.src.modules.writer_prompt import SUMMARY_CONFIRMATION, SUMMARY_ACTION, CHECK_INTENT, CLARIFY_INTENT, INTENT_CONFIRMATION, CATEGORY

from utils import get_mongodb_client, get_redis_client, get_minio_client

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')

logger = logging.getLogger(__name__)
CONVERSATION_DOC_KEY = "conversation:{file_id}"


def reset_or_add(left: int, right: int | str):
    if isinstance(right, str):
        return int(right.split("RESET_")[-1])
    return operator.add(left or 0, right or 0)


async def _get_summary(state: dict, prompt: str, input_context: str = "") -> tuple[str, int, int]:
    llm = get_llm(state["model_id"], state["llm_key"], state["max_tokens"])
    short_answer_prompt = "\nKeep the generated content clear and concise." if state["short_answer"] else ""
    content = f"""Your writing note:
    {state["dependencies"]}{short_answer_prompt}
    Requested output document type/template:
    {state.get("template_key", "")}
    Generate summary in {state["language"]}
    """
    error, success, summary, input_tokens, output_tokens = await get_answer(state["_id"], llm, prompt + "\n" + content + input_context)
    if not success:
        raise error
    return summary, input_tokens, output_tokens


async def _get_summary_status(state: dict, answer: str):
    llm = get_llm(state["model_id"], state["llm_key"], state["max_tokens"])
    content = f"""Chat history:
    {state["chat_history"]}
    Current Plan:
    {state["dependencies"]}
    User's answer:
    {answer}
    """
    return await get_answer_with_schema(state["_id"], llm, SUMMARY_CONFIRMATION, content, SummaryStatus)


async def _get_summary_action(state: dict, answer: dict):
    llm = get_llm(answer["model_id"], answer["llm_key"], state["max_tokens"])
    content = f"""Chat history:
    {state["chat_history"]}
    Current plan:
    {state["dependencies"]}
    Current template key:
    {state.get("template_key", "")}
    User's latest message:
    {answer["question"]}
    """
    return await get_answer_with_schema(state["_id"], llm, SUMMARY_ACTION, content, SummaryAction)


def _is_ask_clarify(state: dict) -> bool:
    for key, value in state["dependencies"].model_dump().items():
        if isinstance(value, list) or isinstance(value, str):
            if not len(value):
                return True
        elif isinstance(value, int):
            # Zero is not a meaningful report length, but remains valid for other integer fields.
            if key == "word_count" and value <= 0:
                return True
        else:
            # For other types, check truthiness
            if not value:
                return True
    return False


async def _get_clarify_question(
    state: dict,
    remaining_keys: list[str],
    prompt: str,
    input_context: str = "",
    include_summary: bool = True,
):
    llm = get_llm(state["model_id"], state["llm_key"], state["max_tokens"])
    short_answer_prompt = "\nKeep the generated content clear and concise." if state["short_answer"] else ""
    summary = ""
    sum_error = None
    sum_success = True
    sum_input_tokens = 0
    sum_output_tokens = 0
    if include_summary:
        content = f"""Current chat history:
        {state["chat_history"]}
        Note:
        {state["dependencies"]}
        Generate summary in {state["language"]}{short_answer_prompt}
        """
        summary_prompt = """You will be given a chat history and a research note
        Summarize the research note to let the user know what they currently have
        **Your response must be ONLY the summary itself.**
        """
        sum_error, sum_success, summary, sum_input_tokens, sum_output_tokens = await get_answer(
            state["_id"], llm, content + input_context, summary_prompt
        )
    final_error = None
    if not sum_success:
        final_error = sum_error

    # Format the remaining keys as a numbered list with explicit instructions
    remaining_keys_list = "\n".join([f"{i+1}. {key.replace('_', ' ')}" for i, key in enumerate(remaining_keys)])

    content = f"""Current chat history:
    {state["chat_history"]}
    Your note:
    {state["dependencies"]}
    User's current intent:
    {CATEGORY[state["category"] - 1]}

    Instructions:
    You must ask exactly {len(remaining_keys)} questions - one question for each of the following fields:
    {remaining_keys_list}

    Rules:
    1. You must generate exactly {len(remaining_keys)} numbered questions
    2. Each question must correspond to one field from the list above
    3. Do not skip any fields from the list
    4. Do not combine multiple fields into one question
    5. Each question should include example answers or suggestions to help the user

    Generate all {len(remaining_keys)} questions in {state["language"]}{short_answer_prompt}
    """
    error, success, questions, input_tokens, output_tokens = await get_answer(state["_id"], llm, content + input_context, prompt)
    if not success:
        final_error = error
    question_output = f"{summary}\n\n{questions}" if include_summary else questions
    return final_error, success and sum_success, question_output, sum_input_tokens + input_tokens, sum_output_tokens + output_tokens


def _is_ask_intent(state: dict) -> bool:
    return state["intent"].change_intent


async def _check_intent(state: dict, answer: dict) -> Intent:
    llm = get_llm(answer["model_id"], answer["llm_key"], state["max_tokens"])
    content = f"""Chat history:
    {state["chat_history"]}
    User's query:
    {answer["question"]}
    User's current intent:
    {CATEGORY[state["category"] - 1]}
    """
    error, success, intent, input_tokens, output_tokens = await get_answer_with_schema(state["_id"], llm, CHECK_INTENT, content, Intent)
    if not success:
        if error.status_code in [401, 403, 429, 500]:
            raise error
        intent = Intent(change_intent=False, new_intent=0)
    return intent, input_tokens, output_tokens


async def _get_clarify_intent_question(state: dict):
    short_answer_prompt = "\nKeep the generated content clear and concise." if state["short_answer"] else ""
    content = f"""User's current intent:
    {state["category"]}
    User's want to change intent to:
    {CATEGORY[state["intent"].new_intent - 1]}
    Generate question in {state["language"]}{short_answer_prompt}
    """
    llm = get_llm(state["model_id"], state["llm_key"], state["max_tokens"])
    return await get_answer(state["_id"], llm, CLARIFY_INTENT + "\n" + content)


async def _get_intent_status(state: dict, answer: dict):
    llm = get_llm(answer["model_id"], answer["llm_key"], state["max_tokens"])
    content = f"""Chat history:
    {state["chat_history"]}
    User's answer:
    {answer["question"]}
    """
    return await get_answer_with_schema(state["_id"], llm, INTENT_CONFIRMATION, content, IntentStatus)


class DataHelper:
    def __init__(
        self, 
        user_id: str, 
        model_id: str, 
        llm_key: str, 
        conv_id: str, 
        max_tokens: int = 10000, 
        mongo_client: AsyncMongoClient | None = None
    ):
        self.model_id = model_id
        self.llm = get_llm(self.model_id, llm_key, max_tokens)
        self.minioClient = get_minio_client()
        self.conv_id = conv_id
        self.user_id = user_id
        self.mongo_client = mongo_client if mongo_client else get_mongodb_client()
        db = self.mongo_client["chatbot"]
        self.collection = db[self.conv_id]
        self.redis_client = get_redis_client()
        self.input_tokens = 0
        self.output_tokens = 0

    @property
    def id(self):
        return self.conv_id[:8]

    def __str__(self):
        return self.id
    
    async def _get_text_from_image(self, file_path: Path) -> str:
        """
        Encodes an image and uses the multimodal LLM to extract its content as text.
        """
        logger.info(f"[{self}] Processing image file with LLM: {file_path.name}")
        with open(file_path, "rb") as image_file:
            image_bytes = image_file.read()

        base64_image = base64.b64encode(image_bytes).decode("utf-8")
        extension = file_path.suffix.lower().lstrip('.')
        mime_type = f"image/{'jpeg' if extension == 'jpg' else extension}"

        prompt = {
            "role": "user",
            "content": [
                {
                    "type": "text",
                    "text": (
                        "You are an expert data extractor. "
                        "Analyze this image and extract all text and relevant information. "
                        "If it contains text, transcribe the text exactly as you see it. "
                        "If it is a chart or graph, describe its key findings, data points, and labels."
                    ),
                },
                {
                    "type": "image",
                    "source_type": "base64",
                    "data": base64_image,
                    "mime_type": mime_type,
                },
            ]
        }
        
        error, success, response, input_tokens, output_tokens = await get_answer(self.id, self.llm, prompt)
        if not success:
            raise error
        return response, input_tokens, output_tokens

    async def _get_text_from_tabular(self, file_path: Path) -> str:
        """
        Reads a tabular file and uses an LLM to provide a comprehensive summary.
        """
        logger.info(f"[{self}] Processing tabular file with LLM: {file_path.name}")
        if file_path.suffix.lower() == ".csv":
            df = pd.read_csv(file_path)
        elif file_path.suffix.lower() in [".xlsx", ".xls"]:
            df = pd.read_excel(file_path)
        else:
            raise AIERROR(602, f"Invalid file type: {file_path.suffix}")

        table_string = df.to_string()

        prompt = (
            "You are an expert data analyst. The following text is the content of a spreadsheet.\n"
            "Analyze it and provide a comprehensive text summary, including a general description, "
            "key columns, and any notable patterns or trends.\n\n"
            f"--- SPREADSHEET DATA ---\n{table_string}"
        )
        error, success, response, input_tokens, output_tokens = await get_answer(self.id, self.llm, prompt)
        if not success:
            raise error
        return table_string + "\n" + response, input_tokens, output_tokens

    async def _get_text_from_word(self, file_path: Path) -> str:
        logger.info(f"[{self}] Processing word file: {file_path.name}")
        if file_path.suffix.lower() == ".docx":
            doc = docx.Document(file_path)
            content = []
            for docpara in doc.paragraphs:
                content.append(docpara.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            content.append(paragraph.text)
            return "\n\n".join(content)
        else:
            raise AIERROR(602, f"Invalid file type: {file_path.suffix}")

    async def _load_file_content(self, url: str) -> str:
        input_tokens, output_tokens = 0, 0
        object_name = url.split("users/")[-1]    
        redis_key = CONVERSATION_DOC_KEY.format(file_id=object_name)
        result = await self.redis_client.get(redis_key)
        if result:
            return json.loads(result), input_tokens, output_tokens
        else:
            result = await self.collection.find_one({"_id": object_name})
            if result:
                await self.redis_client.set(redis_key, json.dumps(result["content"]), ex=settings.REDIS_TTL)
                return result["content"], input_tokens, output_tokens
            else:
                file_path = f"{settings.TMP_FOLDER}/{object_name.split("/")[-1]}"
                os.makedirs(settings.TMP_FOLDER, exist_ok=True)
                self.minioClient.fget_object("users", object_name, file_path)
                path = Path(file_path)
                extension = path.suffix.lower()
                if extension == ".pdf":
                    logger.info(f"[{self}] Reading pdf file: {path.name}")
                    loader = PyPDFLoader(file_path)
                    doc = await loader.aload() 
                    content = '\n\n'.join([page.page_content for page in doc])
                elif extension in [".txt", ".md"]:
                    logger.info(f"[{self}] Reading text file: {path.name}")
                    with open(path, "r", encoding="utf-8", errors="ignore") as f:
                        content = f.read()
                elif extension in [".docx"]:
                    content = await self._get_text_from_word(path)
                elif extension in [".csv", ".xlsx", ".xls"]:
                    content, input_tokens, output_tokens = await self._get_text_from_tabular(path)
                elif extension in [".png", ".jpg", ".jpeg"]:
                    content, input_tokens, output_tokens = await self._get_text_from_image(path)
                else:
                    logger.info(f"[{self}] Unsupported file type: {extension} for file {path.name}")
                    content = f"Error: Unsupported file type '{extension}' for file {path.name}."
                os.remove(file_path)
                await self.redis_client.set(redis_key, json.dumps(content), ex=settings.REDIS_TTL)
                await self.collection.insert_one({"_id": object_name, "content": content})
                return content, input_tokens, output_tokens
            
    async def load_file_content(self):
        try:
            prefix = f"{self.user_id}/{settings.MINIO_PREFIX}/{self.conv_id}"
            tasks = [
                self._load_file_content(obj.object_name) 
                for obj in self.minioClient.list_objects("users", prefix=prefix, recursive=True)
            ]
            files_content = await asyncio.gather(*tasks)
            results = [result[0] for result in files_content]
            input_tokens = [result[1] for result in files_content]
            output_tokens = [result[2] for result in files_content]
            self.input_tokens += sum(input_tokens)
            self.output_tokens += sum(output_tokens)
            input_context = "\nInput files:\n" + "\n\n".join(results)
        except Exception:
            input_context = ""
        finally:
            return input_context, self.input_tokens, self.output_tokens


def tool_matching(query: str, tools: list[str]) -> list[str]:
    if query in tools:
        return query
    else:
        best_match = rapidfuzz.process.extractOne(query, tools, scorer=rapidfuzz.fuzz.WRatio)
        return best_match[0]
